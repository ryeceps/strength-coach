"""Check published report prescriptions against their workbook sources (openpyxl)."""
import json
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]


def report_rows(name):
    plan = json.loads((root / "examples" / f"{name}.json").read_text(encoding="utf-8"))
    return [(week["number"], day["title"], item["name"], item["prescription"])
            for mesocycle in plan["mesocycles"] for week in mesocycle["weeks"]
            for day in week["days"] for item in day["exercises"]]


rebuild = load_workbook(root / "assets/generated-programs/12_week_strongman_program.xlsx", read_only=True, data_only=True)
source = [(row[0], row[3], row[5], row[6]) for row in list(rebuild["Program"].values)[3:] if isinstance(row[0], int)]
assert report_rows("strongman-rebuild") == source, "Rebuild report diverges from original workbook"

pyramid = load_workbook(root / "assets/generated-programs/12_week_strongman_4_day_rpe_pyramid_program.xlsx", read_only=True, data_only=True)
source = [(row[0], row[4], row[7], row[8]) for row in list(pyramid["Plan Detail"].values)[1:] if isinstance(row[0], int)]
assert report_rows("strongman-rpe-pyramid") == source, "Pyramid report diverges from original workbook"

bodybuilding = load_workbook(root / "site/assets/generated-programs/bodybuilding-priority-block.xlsx", read_only=True, data_only=True)
source = [(row[1], row[4], row[5], row[6]) for row in list(bodybuilding["Program"].values)[3:] if isinstance(row[1], int)]
assert report_rows("bodybuilding-priority-block") == source, "Bodybuilding workbook diverges from report JSON"

print("Source fidelity passed: two original strongman workbooks and the new bodybuilding workbook")
