"""One-time, repeatable conversion of the two existing public workbooks to report JSON.

Requires openpyxl. Prescriptions and notes are copied from the workbook cells;
sets/RPE and muscle tags are coaching metadata used by the report analysis.
"""
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BOOKS = ROOT / "assets" / "generated-programs"
OUT = ROOT / "examples"
OUT.mkdir(exist_ok=True)


def mapping(name):
    n = name.lower()
    if "hamstring" in n and "dead bug" in n:
        return ["hamstrings", "abs"], ["glutes"]
    if "dead bug" in n:
        return ["abs"], []
    if "hamstring" in n and "deadlift" not in n:
        return ["hamstrings"], ["glutes"]
    if "deadlift" in n or "romanian" in n:
        return ["glutes", "hamstrings", "spinal-erectors"], ["upper-back", "forearms"]
    if "squat" in n:
        return ["quads", "glutes"], ["adductors", "spinal-erectors"]
    if "farmer" in n or "carry" in n:
        return ["forearms", "upper-back"], ["abs", "glutes"]
    if "log" in n or "overhead" in n:
        return ["front-delts", "triceps"], ["upper-back", "abs"]
    if "bench" in n and ("row" in n or "pulldown" in n):
        return ["chest", "triceps", "upper-back", "lats"], ["front-delts", "biceps"]
    if "bench" in n:
        return ["chest", "triceps"], ["front-delts"]
    if "row" in n or "pulldown" in n:
        return ["upper-back", "lats"], ["biceps", "rear-delts"]
    if "shrug" in n:
        return ["upper-back", "rear-delts"], ["forearms"]
    if "rear" in n and "triceps" in n:
        return ["rear-delts", "triceps"], ["upper-back"]
    if "pressdown" in n and "rear" in n:
        return ["triceps", "rear-delts"], ["upper-back"]
    if "face pull" in n or "rear" in n:
        return ["rear-delts", "upper-back"], []
    if "lateral" in n:
        return ["side-delts", "biceps"], []
    if "curl" in n:
        return ["biceps"], ["forearms"]
    if "pressdown" in n or "triceps" in n:
        return ["triceps"], []
    return ["abs"], []


def metadata(prescription):
    text = str(prescription)
    if "top 1 x" in text.lower() and "backoff" in text.lower():
        match = re.search(r"backoff\s+(\d+)\s*x", text, re.I)
        sets = 1 + (int(match.group(1)) if match else 0)
    elif "Top 3" in text:
        sets = 4
    else:
        match = re.search(r"(\d+)\s*(?:-\s*\d+)?\s*(?:x|rounds)", text, re.I)
        sets = int(match.group(1)) if match else 2
    rpes = []
    for match in re.finditer(r"RPE\s*(\d+(?:\.\d+)?)(?:\s*[-–]\s*(\d+(?:\.\d+)?))?", text, re.I):
        rpes.append(float(match.group(2) or match.group(1)))
    return sets, max(rpes) if rpes else None


def reason_for(name):
    n = str(name).lower()
    if "log clean and press" in n:
        return "Practice the clean and press with the contest implement."
    if "clean once" in n:
        return "Get more press work without another full clean on every rep."
    if "deadlift" in n and "romanian" not in n:
        return "Pull heavy. Keep the deadlift moving."
    if "front squat" in n:
        return "Build the quads and stay upright under the bar."
    if "back squat" in n:
        return "Make the squat a main lift and track it each week."
    if "farmer" in n:
        return "Carry heavy. Train grip and bracing under load."
    if "hamstring" in n or "romanian" in n:
        return "Train the hamstrings without another heavy pull from the floor."
    if "bench" in n and ("row" in n or "pulldown" in n):
        return "Press and row in the same slot. Keep upper-back work in the week."
    if "slingshot" in n:
        return "Add close-grip press work after the upper-back work."
    if "bench" in n:
        return "Build the bench without taking work away from the event lift."
    if "row" in n or "pulldown" in n:
        return "Build the upper back for pressing and carries."
    if "face pull" in n or "rear" in n:
        return "Train the rear delts. Cut this first if time is short."
    if "shrug" in n:
        return "Train traps and grip after the main work."
    if "lateral" in n or "curl" in n or "pump" in n:
        return "Finish with a small dose of arm or shoulder work."
    return "Support the main lift without adding much fatigue."


def technique_for(name):
    n = str(name).lower()
    if "log" in n:
        return "Brace before the press and keep the implement close through the clean.", ["Axle clean and press when the log is unavailable"]
    if "deadlift" in n and "romanian" not in n:
        return "Set the brace and keep the bar path close; stop when positions deteriorate.", ["Trap-bar deadlift when the setup or tolerance requires it"]
    if "front squat" in n:
        return "Keep the elbows high and torso stacked over the midfoot.", ["Safety-bar squat when front rack mobility limits the lift"]
    if "back squat" in n:
        return "Use the normal stance and depth; maintain the brace out of the bottom.", ["Safety-bar squat if the bar position is not tolerated"]
    if "farmer" in n:
        return "Stand tall, keep the steps controlled, and set the handles down before grip or posture fails.", ["Trap-bar carry if farmer handles are unavailable"]
    if "low-cable row" in n:
        return "Brace the feet on a fixed support and pull without rocking the torso.", ["Supported one-arm cable row if no stable foot brace exists"]
    if "slingshot" in n:
        return "Judge effort on the assisted lift rather than comparing its load with raw bench.", []
    return "", []


def exercise(name, prescription, note=""):
    primary, secondary = mapping(str(name))
    sets, rpe = metadata(prescription)
    cue, alternatives = technique_for(name)
    item = {"name": str(name), "prescription": str(prescription), "sets": sets,
            "primaryMuscles": primary, "secondaryMuscles": secondary,
            "reason": reason_for(name)}
    if str(name).lower().startswith("optional"):
        item["optional"] = True
    if rpe is not None:
        item["rpe"] = rpe
    if note:
        item["note"] = str(note)
    if cue:
        item["cue"] = cue
    if alternatives:
        item["alternatives"] = alternatives
    return item


def mesocycles(weeks, phases):
    output = []
    for title, start, end, rationale in phases:
        selected = [weeks[number] for number in range(start, end + 1)]
        output.append({"title": title, "focus": title, "rationale": rationale, "weeks": selected})
    return output


def rebuild():
    workbook = load_workbook(BOOKS / "12_week_strongman_program.xlsx", read_only=True, data_only=True)
    rows = list(workbook["Program"].values)[3:]
    grouped = defaultdict(lambda: defaultdict(list))
    meta = {}
    for row in rows:
        if len(row) < 7 or not isinstance(row[0], int):
            continue
        week, phase, day, focus, slot, name, prescription = row[:7]
        grouped[week][day].append(exercise(name, prescription))
        meta[(week, day)] = (phase, focus)
    weeks = {}
    for number, days in sorted(grouped.items()):
        sessions = []
        for day, exercises in sorted(days.items()):
            phase, focus = meta[(number, day)]
            sessions.append({"weekday": [0, 2, 4][day - 1], "title": focus,
                             "focus": "Four capped exercise slots, with the priority work first.", "exercises": exercises})
        weeks[number] = {"number": number, "phase": meta[(number, 1)][0],
                         "focus": meta[(number, 1)][0], "note": "Use RPE to adjust the day's load; keep heavy work technically clean.",
                         "days": sessions}
    return {"schemaVersion": 1, "id": "strongman-rebuild", "title": "12-week strongman rebuild",
            "sport": "Strongman", "objective": "Rebuild into heavy triples",
            "summary": "A three-day block centered on log press, deadlift, front squat, and farmer carries, with bench and back work twice each week.",
            "rationale": "Rebuild tolerance first, raise strength through the middle weeks, then use controlled heavy triples. Main events retain priority; accessories stay capped.",
            "goals": ["Rebuild repeatable log, pull, squat, and carry work.", "Keep bench and upper-back support twice weekly.", "Reach heavy triples without singles or max attempts."],
            "priorityMuscles": ["quads", "hamstrings", "upper-back"], "maxExercisesPerDay": 4,
            "assumptions": ["Three nonconsecutive training days are available each week.", "No event date or personal recovery log was supplied; adjust with actual performance.", "Combined slots may train more than one muscle; set counts are coaching approximations."],
            "sources": [{"label": "Original 12-week workbook", "url": "https://github.com/ryeceps/strength-coach/blob/main/assets/generated-programs/12_week_strongman_program.xlsx"}],
            "mesocycles": mesocycles(weeks, [
                ("Rebuild", 1, 4, "Submaximal work establishes rhythm and tolerance."),
                ("Strength build", 5, 8, "Load and effort rise while volume stays controllable."),
                ("Heavy triples and pivot", 9, 12, "Practice heavier triples, then deload or take one controlled test only when ready.")])}


def pyramid():
    workbook = load_workbook(BOOKS / "12_week_strongman_4_day_rpe_pyramid_program.xlsx", read_only=True, data_only=True)
    rows = list(workbook["Plan Detail"].values)[1:]
    grouped = defaultdict(lambda: defaultdict(list))
    meta = {}
    for row in rows:
        if len(row) < 10 or not isinstance(row[0], int):
            continue
        week, phase, intensity, day, focus, slot, block, name, prescription, note = row[:10]
        grouped[week][day].append(exercise(name, prescription, note))
        meta[(week, day)] = (phase, intensity, focus)
    weeks = {}
    for number, days in sorted(grouped.items()):
        sessions = []
        for day, exercises in sorted(days.items()):
            phase, intensity, focus = meta[(number, day)]
            sessions.append({"weekday": [0, 1, 3, 5][day - 1], "title": focus,
                             "focus": "Focused lift, support, then accessory; cut accessory work first if time runs long.",
                             "exercises": exercises})
        phase, intensity, _ = meta[(number, 1)]
        weeks[number] = {"number": number, "phase": phase, "focus": intensity,
                         "note": "Top sets are RPE-capped. Move to the backoff work without chasing a missed number.",
                         "days": sessions}
    return {"schemaVersion": 1, "id": "strongman-rpe-pyramid", "title": "12-week strongman RPE pyramid",
            "sport": "Strongman", "objective": "Home-gym strength build",
            "summary": "A four-day home-gym mesocycle with log, deadlift, squat, upper-back work, and rotating six-, five-, and three-rep top sets.",
            "rationale": "Three-week RPE waves create varied exposures, while pivot weeks reduce accumulated work. Raw bench follows log, and the upper-back day uses stable cable work.",
            "goals": ["Keep log, deadlift, and back squat as the measurable priorities.", "Vary top-set reps without singles.", "Protect the lower back and fit sessions into 45–60 minutes."],
            "priorityMuscles": ["upper-back", "hamstrings", "quads"], "maxExercisesPerDay": 4,
            "assumptions": ["This is a general strongman block without a dated contest or required carry event.", "Equipment includes a log, barbell, rack, bench, Slingshot, and stable low-cable row setup.", "The top set and backoffs count as working sets; warm-up ramps do not."],
            "sources": [{"label": "Detailed source template", "url": "https://github.com/ryeceps/strength-coach/blob/main/templates/strongman/strongman-rpe-pyramid-four-day.md"},
                        {"label": "Original 12-week workbook", "url": "https://github.com/ryeceps/strength-coach/blob/main/assets/generated-programs/12_week_strongman_4_day_rpe_pyramid_program.xlsx"}],
            "mesocycles": mesocycles(weeks, [
                ("Volume and rhythm", 1, 4, "Six-, five-, and three-rep exposures followed by a lower-volume pivot."),
                ("Build and vary", 5, 8, "Repeat the wave with progression only when recovery and technique support it."),
                ("Specific work and deload", 9, 12, "Heavy triples stay controlled; week 12 is a deload or one optional controlled test.")])}


for name, program in [("strongman-rebuild", rebuild()), ("strongman-rpe-pyramid", pyramid())]:
    path = OUT / f"{name}.json"
    path.write_text(json.dumps(program, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"Wrote {path}")

